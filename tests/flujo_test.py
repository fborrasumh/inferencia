"""Flujo completo de InferencIA en el navegador, con IA simulada que MIENTE (cita inventada, nivel inflado, solución filtrada):
   portada, idiomas, ejemplo sin clave, práctica de las 19 habilidades en los tres idiomas, laboratorio con los seis diseños,
   script R, croquis, comentarios de la IA al informe, progreso, evidencia con huella, Word, CSV, GIFT, profesorado, persistencia
   y móvil.   Uso: python3 tests/flujo_test.py"""
import asyncio, sys, pathlib, json, csv, io, re
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from forja_test import Harness
ROOT = pathlib.Path(__file__).resolve().parent.parent
MODO = {'fuga': True}

async def reply(prov, system, user, info):
    if 'Comenta la interpretación' in user:
        en = 'English' in user
        return ('j', {'criterios': [
            {'id': 'hipotesis', 'nivel': 'logrado', 'cita': 'at least one differs' if en else 'al menos una difiere', 'sug': 'Bien.'},
            {'id': 'anava', 'nivel': 'logrado', 'cita': 'esta frase no está en el texto del estudiante de ninguna manera', 'sug': 'Añade el CV del 7,77 %.'},
            {'id': 'supuestos', 'nivel': 'parcial', 'cita': 'Shapiro-Wilk', 'sug': 'Informa Levene.'},
            {'id': 'comparaciones', 'nivel': 'logrado', 'cita': 'Capiro', 'sug': 'Bien.'},
            {'id': 'agronomica', 'nivel': 'ausente', 'cita': '', 'sug': 'Añade limitaciones.'}]})
    if 'Explica en un máximo' in user:
        sol = re.search(r'SOLUCION=([-\d.]+)', MODO.get('sol', '')) if MODO.get('sol') else None
        frase_fuga = f" El resultado correcto es {MODO['sol']}." if MODO['fuga'] and MODO.get('sol') else ''
        return ('t', 'Revisa los grados de libertad que has usado.' + frase_fuga + ' Vuelve a mirar la tabla.')
    return ('t', 'ok')

async def main():
    ok = True
    def check(nombre, cond, extra=''):
        nonlocal ok; ok &= bool(cond); print(('✓' if cond else '✕'), nombre, ('· ' + str(extra)[:300]) if extra and not cond else '')
    async with Harness(ROOT/'index.html', reply, init_scripts=["localStorage.setItem('ia_openai_key','sk-test-clave-larga-123')"]) as h:
        p = h.page
        check('carga sin errores de página', h.errors == [], h.errors)
        check('CSP bloquea otros orígenes', await h.csp_bloquea())
        for prov, esperado in (('openai', 'gpt-5-mini'), ('gemini', 'gemini-2.5-pro'), ('anthropic', 'claude-sonnet-4-5')):
            ops = await h.modelos(prov); check(f'desplegable de modelos ({prov})', esperado in ops and 'whisper-1' not in ops, ops)
        await p.select_option('#selProv', 'openai'); await p.keyboard.press('Escape')
        check('no faltan traducciones (HTML, pasos, botones, glosario)', await h.faltan_traducciones() == [], await h.faltan_traducciones())
        # registro de traducciones que faltan en tiempo de ejecución (todo lo que pasa por t())
        await p.evaluate("""window._faltan=new Set();const _t=window.t;window.t=function(k,...a){if(LANG!=='es'&&typeof k==='string'&&/\\p{L}{3}/u.test(k)&&!(I18N[LANG]&&k in I18N[LANG])&&!(I18N_BASE[LANG]&&k in I18N_BASE[LANG]))window._faltan.add(k);return _t(k,...a)}""")
        # ── ejemplo sin clave en los tres idiomas ──
        for l, palabra in (('es', 'variedades de papa'), ('en', 'potato varieties'), ('pt', 'variedades de batata')):
            await h.idioma(l); await p.evaluate('go(0)'); await p.click('#btnEjemplo'); await p.wait_for_timeout(500)
            fbok = await p.evaluate("S.inf.fb && S.inf.fb.crit.every(c=>c.ok)")
            check(f'ejemplo ({l}): análisis DBCA, interpretación con citas verificadas y cifras coherentes', await p.evaluate("step===2 && S.lab.res && !S.lab.res.error") and palabra in (await p.evaluate("S.inf.texto")) and fbok and await p.locator('#fbInf .okbox').count() == 1, await p.inner_text('#fbInf'))
            check(f'ejemplo ({l}): tabla ANAVAR, letras, supuestos y script R', await p.locator('.ag-anova').count() >= 1 and await p.locator('.ag-letra').count() == 4 and 'aov(' in await p.inner_text('#scriptR'))
            await p.evaluate('go(3)'); check(f'ejemplo ({l}): progreso con 7 barras de dominio', await p.locator('.ag-meter').count() == 7)
        await h.idioma('es')
        r = await p.evaluate("S.lab.res.tabla.map(f=>[f.fuente,f.gl,+((f.F||0).toFixed(2))])")
        check('ejemplo: F de tratamientos = 23,82 con 3 y 9 gl (como dice el texto del ejemplo)', r[1] == ['Tratamientos', 3, 23.82] and r[2][1] == 9, r)
        # ── práctica: las 19 habilidades en los tres idiomas; respuesta correcta, error típico y solución resuelta ──
        await p.evaluate("S=nuevoS();guarda();go(1)"); await p.fill('#alias', 'prueba-01'); await p.click('#btnAlias'); await p.wait_for_timeout(200)
        skills = await p.evaluate("SKILLS.map(s=>s.id)")
        for l in ('es', 'en', 'pt'):
            await h.idioma(l); malos = []
            for sk in skills:
                await p.evaluate(f"nuevaActividad('{sk}','elegida');pintaPractica()")
                txt = await p.inner_text('#cardEj')
                if len(txt) < 80 or 'undefined' in txt or 'NaN' in txt: malos.append(sk)
            check(f'las {len(skills)} habilidades se pintan en {l} (sin undefined/NaN)', not malos and h.errors == [], [malos, h.errors])
        await h.idioma('es')
        resueltas, diagnos = [], []
        for sk in skills:
            await p.evaluate(f"nuevaActividad('{sk}','elegida');pintaPractica()")
            # primero un error típico (si la habilidad lo tiene), después la respuesta correcta
            mut = await p.evaluate("(()=>{const e=ej();const q=e.q.find(q=>q.mut&&q.mut.some(m=>Number.isFinite(m.v)&&Math.abs(m.v-q.sol)>Math.max(q.tol.abs||0,(q.tol.rel||0)*Math.abs(q.sol))*3));if(!q)return null;const m=q.mut.find(m=>Number.isFinite(m.v)&&Math.abs(m.v-q.sol)>Math.max(q.tol.abs||0,(q.tol.rel||0)*Math.abs(q.sol))*3);return {id:q.id,v:m.v.toFixed(q.dec).replace('.',','),k:m.k}})()")
            async def rellena(sobre=None):
                qs = await p.evaluate("ej().q.map(q=>({id:q.id,tipo:q.tipo||'num',sol:q.tipo==='opc'?q.sol:q.sol.toFixed(q.dec).replace('.',',')}))")
                for q in qs:
                    v = sobre['v'] if sobre and sobre['id'] == q['id'] else q['sol']
                    if q['tipo'] == 'opc': await p.check(f"input[name='r_{q['id']}'][value='{q['sol']}']")
                    else: await p.fill(f"#r_{q['id']}", v)
                await p.click('#btnComprobar'); await p.wait_for_timeout(120)
            if mut:
                await rellena(mut); fb = await p.inner_text('#fbEj')
                diag = await p.evaluate(f"diagTxt('{sk}','{mut['k']}')")
                if diag[:25] in fb: diagnos.append(sk)
            await rellena()
            if await p.evaluate("S.act.fin && S.act.okFinal>0") and await p.locator('.ag-sol').count(): resueltas.append(sk)
        check('respuesta correcta aceptada y solución resuelta visible en las 19 habilidades', len(resueltas) == len(skills), set(skills) - set(resueltas))
        con_mut = await p.evaluate("SKILLS.filter(s=>ejercicio('prueba-01',s.id,0).q.some(q=>q.mut&&q.mut.length)).length")
        check('el error típico se diagnostica con su explicación', len(diagnos) >= con_mut - 1, [len(diagnos), con_mut])
        check('el progreso registra los intentos', await p.evaluate("S.eventos.length") == len(skills))
        # ── tutor IA que intenta soplar la solución: el código elimina la frase ──
        await p.evaluate("nuevaActividad('mds','elegida');pintaPractica()")
        sol = await p.evaluate("(()=>{const q=ej().q[0];return q.sol.toFixed(3).replace('.',',')})()"); MODO['sol'] = sol
        await p.fill('#r_dms', '1'); await p.fill('#r_nsig', '0'); await p.click('#btnComprobar'); await p.wait_for_timeout(150)
        await p.click('#btnExplica'); await p.wait_for_selector('#mEnvio.on', timeout=5000)
        check('aviso previo al envío con muestra', 'EJERCICIO' in await p.inner_text('#envioTxt'))
        await p.check('#envioOk'); await p.click('#envioSi'); await h.esp("S.act.ia")
        ia = await p.evaluate("S.act.ia")
        check('la frase con la solución se elimina de la explicación de la IA', ia['quitadas'] == 1 and sol not in ia['texto'], ia)
        check('la petición al tutor no incluye la solución', sol not in h.seen[-1]['user'] and 'español' in h.seen[-1]['user'])
        # ── laboratorio: los seis diseños con sus datos de ejemplo ──
        await p.evaluate('go(2)')
        for dis in ('dca', 'dbca', 'cl', 'factorial', 'crossover', 'regresion'):
            await p.select_option('#labDis', dis); await p.click('#btnDatosEj'); await p.wait_for_timeout(150)
            for met in (['tukey', 'duncan', 'mds', 'dunnett'] if dis in ('dca', 'dbca', 'cl') else ['tukey']):
                if not await p.is_disabled('#labMet'): await p.select_option('#labMet', met)
                await p.click('#btnLab'); await h.esp("S.lab.res"); await p.wait_for_timeout(200)
                res = await p.evaluate("S.lab.res && !S.lab.res.error"); txt = await p.inner_text('#resLab')
                check(f'laboratorio {dis} · {met}: análisis sin errores', res and 'NaN' not in txt and 'undefined' not in txt and h.errors == [], [await p.evaluate("S.lab.res && S.lab.res.error"), h.errors])
        await p.select_option('#labDis', 'dca'); await p.click('#btnDatosEj'); await p.click('#btnLab'); await h.esp("S.lab.res && S.lab.res.diseno==='dca'")
        await h.descarga('#btnDlR', '/tmp/inf.R'); await h.descarga('#btnDlDatos', '/tmp/datos.csv')
        r_txt = pathlib.Path('/tmp/inf.R').read_text(encoding='utf-8')
        check('script R descargable con aov, shapiro y prueba de comparaciones', 'aov(' in r_txt and 'shapiro.test' in r_txt and 'HSD.test' in r_txt)
        check('datos.csv descargable', pathlib.Path('/tmp/datos.csv').read_text(encoding='utf-8-sig').startswith('tratamiento;'))
        # errores de entrada: mensaje claro, no excepción
        await p.fill('#csvLab', 'a;b\n1;x'); await p.wait_for_timeout(500); await p.click('#btnLab'); await p.wait_for_timeout(300)
        check('datos inválidos: aviso claro, sin excepción', await p.locator('#resLab .aviso').count() == 1 and h.errors == [])
        # croquis
        await p.select_option('#crDis', 'cl'); await p.fill('#crTrt', 'A, B, C, D, E'); await p.click('#btnCroq')
        celdas = await p.evaluate("S.croq.filas"); check('croquis de cuadrado latino válido (cada tratamiento una vez por fila y columna)', all(len(set(f)) == 5 for f in celdas) and all(len({f[j] for f in celdas}) == 5 for j in range(5)))
        # ── interpretación con IA que miente: cita inventada → «parcial»; cifra que no está en el análisis → aviso ──
        await p.select_option('#labDis', 'dbca'); await p.click('#btnDatosEj'); await p.click('#btnLab'); await h.esp("S.lab.res && S.lab.res.diseno==='dbca'")
        inf = await p.evaluate("EJ_INF.es.replace('F = 23,82','F = 25,10')")
        await p.fill('#informe', inf); await p.wait_for_timeout(200); await p.click('#btnInfIA'); await h.esp("S.inf.fb")
        fb = await p.evaluate("S.inf.fb.crit.map(c=>[c.id,c.nivel,c.ok])")
        check('la cita inventada se marca y el «logrado» sin cita verificada baja a «parcial»', fb[1] == ['anava', 'parcial', False] and fb[0] == ['hipotesis', 'logrado', True], fb)
        check('el código detecta la cifra del texto que no está en el análisis (25,10)', '25,1' in await p.inner_text('#fbInf'), await p.inner_text('#fbInf'))
        check('la sugerencia con una cifra ajena (7,77) se señala', '7,77' in ''.join(await p.evaluate("S.inf.fb.sugNuevas")), await p.evaluate("S.inf.fb.sugNuevas"))
        check('la petición del informe lleva la tabla calculada y el idioma', 'RESULTADOS CALCULADOS' in h.seen[-1]['user'] and 'Tratamientos' in h.seen[-1]['user'] and 'español' in h.seen[-1]['user'])
        # ── progreso, evidencia, Word, CSV ──
        await p.evaluate('go(3)'); await p.check('#chkInf')
        await h.descarga('#btnEvid', '/tmp/evid.json'); ev = json.loads(pathlib.Path('/tmp/evid.json').read_text())
        check('evidencia JSON con huella, eventos y laboratorio', ev['formato'] == 'inferencia/1' and len(ev['huella']) == 64 and len(ev['eventos']) >= 19 and ev['laboratorio']['diseno'] == 'dbca' and ev['informe']['texto'])
        await h.descarga('#btnCsv', '/tmp/intentos.csv'); filas = list(csv.reader(io.StringIO(pathlib.Path('/tmp/intentos.csv').read_text(encoding='utf-8-sig')), delimiter=';'))
        check('CSV de intentos (BOM, ;)', len(filas) == len(ev['eventos']) + 1 and filas[0][1] == 'RA')
        h.sin_cdn = True; await h.descarga('#btnWord', '/tmp/inf.docx'); h.sin_cdn = False
        import docx; w = docx.Document('/tmp/inf.docx'); tx = '\n'.join(q.text for q in w.paragraphs); tabs = len(w.tables)
        check('Word sin CDN: se abre con python-docx, con tablas de dominio, ANAVAR e interpretación', 'Informe de práctica autónoma' in tx and tabs >= 4 and 'Planteamos H0' in tx, [tabs, tx[:200]])
        # ── profesorado: la evidencia descargada se verifica; una editada se detecta ──
        ev2 = dict(ev); ev2['dominio'] = dict(ev['dominio']); ev2['dominio']['ic'] = 1.0; pathlib.Path('/tmp/evid_editada.json').write_text(json.dumps(ev2))
        await p.evaluate('PROF=[];go(4)'); await p.set_input_files('#fileEvid', ['/tmp/evid.json', '/tmp/evid_editada.json']); await p.wait_for_timeout(400)
        integ = await p.evaluate("PROF.map(p=>p.integra)")
        check('profesorado: huella válida en la original y ⚠ en la editada', integ == [True, False], integ)
        await p.click('#btnProfEj'); await p.wait_for_timeout(500)
        check('profesorado: grupo de ejemplo con propuestas de mejora', await p.evaluate("PROF.length") == 12 and await p.locator('#resProf li').count() >= 1)
        check('profesorado: autopsia del sílabo con secciones citadas', await p.locator('#autopsia li').count() == 7 and '§' in await p.inner_text('#autopsia'))
        await p.fill('#giftN', '2'); await h.descarga('#btnGift', '/tmp/gift.txt'); g = pathlib.Path('/tmp/gift.txt').read_text(encoding='utf-8')
        check('GIFT para Moodle: 19 categorías, preguntas numéricas con tolerancia y de opción', g.count('$CATEGORY') == 19 and len(re.findall(r'\{#-?[\d.]+:[\d.]+\}', g)) >= 30 and '=Se rechaza' in g or '=No se rechaza' in g)
        falt = await p.evaluate("[...window._faltan]")
        # ── persistencia ──
        await p.reload(); await p.wait_for_timeout(700)
        check('al recargar se conservan el seudónimo, los intentos y el laboratorio', await p.evaluate("S.alias") == 'prueba-01' and await p.evaluate("S.eventos.length") >= 19 and await p.evaluate("S.lab.res.diseno") == 'dbca')
        # error de la IA traducido
        await h.idioma('pt'); await p.evaluate("ENVIO_OK=true;go(2)")
        async def r429(prov, system, user, info): return ('err', 429, '{"error":{"message":"You exceeded your current quota"}}')
        h.reply = r429; await p.click('#btnInfIA'); await p.wait_for_timeout(1500)
        check('un 429 (sin saldo) se explica en el idioma elegido', re.search(r'saldo|crédito|quota|quota', await p.inner_text('#estadoInf'), re.I) and 'ó' not in (await p.inner_text('#estadoInf'))[:0], await p.inner_text('#estadoInf'))
        h.reply = reply; await h.idioma('es')
        check('ninguna frase sin traducir en tiempo de ejecución (en/pt)', not falt, falt)
        check('sin errores de página en todo el flujo', h.errors == [], h.errors)
        # ── móvil: sin desborde en las pantallas con datos, en los tres idiomas ──
        for l in ('es', 'en', 'pt'):
            for pant, js in ((0, ''), (1, "S.alias='movil';nuevaActividad('cl','elegida');"), (2, "aplicaEjemplo();"), (3, "aplicaEjemplo();"), (4, "")):
                d = await h.desborde_movil(f"localStorage.setItem(APP+'_lang','{l}');setLang('{l}');{js}go({pant});" + ("$('btnProfEj').click();" if pant == 4 else ''))
                if d: check(f'móvil {l} pantalla {pant} sin desborde horizontal', False)
        check('móvil (390 px): sin desborde horizontal en las 5 pantallas y 3 idiomas', True)
    print('\n' + ('TODO OK' if ok else 'HAY FALLOS')); return ok

if __name__ == '__main__': sys.exit(0 if asyncio.run(main()) else 1)
